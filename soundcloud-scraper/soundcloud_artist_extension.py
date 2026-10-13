from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.styles import Font

import sc_likes_scraper4_organized_outputs as organizer


def install_artist_index() -> None:
    original_write_workbook = organizer.write_organized_workbook

    def write_workbook_with_artist_index(source_csv_path, cleaned_rows, fieldnames):
        workbook_path = original_write_workbook(source_csv_path, cleaned_rows, fieldnames)
        if not workbook_path:
            return workbook_path

        workbook = load_workbook(workbook_path)
        if "Artist Index" in workbook.sheetnames:
            del workbook["Artist Index"]
        artist_sheet = workbook.create_sheet("Artist Index", 2)
        headers = ["Artist", "Title", "Genre", "SoundCloud URL", "Download URL", "Buy URL"]
        artist_sheet.append(headers)

        sorted_rows = sorted(
            cleaned_rows,
            key=lambda row: (
                str(row.get("artist") or "Unknown Artist").lower(),
                str(row.get("title") or "").lower(),
            ),
        )
        for row in sorted_rows:
            download_url = row.get("final_download_url") or row.get("download_target_url") or ""
            values = [
                row.get("artist") or "Unknown Artist",
                row.get("title") or "",
                row.get("genre") or "",
                row.get("track_url") or "",
                download_url,
                row.get("buy_url") or "",
            ]
            artist_sheet.append(values)
            row_number = artist_sheet.max_row
            for column in (4, 5, 6):
                cell = artist_sheet.cell(row=row_number, column=column)
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = Font(color="0563C1", underline="single")

        artist_sheet.freeze_panes = "A2"
        artist_sheet.auto_filter.ref = artist_sheet.dimensions
        widths = {"A": 28, "B": 42, "C": 24, "D": 48, "E": 48, "F": 48}
        for column, width in widths.items():
            artist_sheet.column_dimensions[column].width = width

        workbook.save(workbook_path)
        workbook.close()
        return workbook_path

    organizer.write_organized_workbook = write_workbook_with_artist_index
