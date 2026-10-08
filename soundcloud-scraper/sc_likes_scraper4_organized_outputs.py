from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
from urllib.parse import urlparse
from datetime import datetime
import csv
import json
import os
import re
import sys
import time

# -----------------------------
# Configuration
# -----------------------------
# Update this for your machine if needed.
BRAVE_PATH = "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser"
USER_DATA_DIR = os.path.abspath("./brave_sc_profile")

BATCH_SIZE = 50
MAX_IDLE_ROUNDS = 20
WAIT_MS = 1800
FAST_SCROLL_WAIT_MS = 900
NORMAL_SCROLL_PIXELS = 7000
FAST_SCROLL_PIXELS = 20000
DUPLICATE_ONLY_THRESHOLD = 5
DEFAULT_GATE_LIMIT = 20
PAGE_TIMEOUT_MS = 90000

CSV_FIELDS = [
    "track_url",
    "title",
    "artist",
    "genre",
    "download_available",
    "download_target_url",
    "download_target_type",
    "buy_url",
    "buy_target_type",
    "discovered_at",
    "source_profile",
    "gate_attempted_at",
    "gate_status",
    "gate_notes",
    "final_download_url",
]


# -----------------------------
# Utilities
# -----------------------------
def utc_now_iso():
    return datetime.utcnow().isoformat(timespec="seconds") + "Z"


def clean_text(value):
    if value is None:
        return None
    value = re.sub(r"\s+", " ", str(value)).strip()
    return value or None


def slugify_profile_name(name: str) -> str:
    value = name.strip().lower()
    value = re.sub(r"[^a-z0-9_-]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    return value or "soundcloud_profile"


def normalize_url(href: str):
    if not href:
        return None

    href = href.strip()
    if href.startswith("/"):
        href = "https://soundcloud.com" + href

    if not href.startswith("https://soundcloud.com/"):
        return None

    href = href.split("?")[0].split("#")[0].rstrip("/")
    return href


def normalize_any_url(href: str):
    if not href:
        return None

    href = href.strip()
    if not href:
        return None

    if href.startswith("//"):
        href = "https:" + href
    elif href.startswith("/"):
        href = "https://soundcloud.com" + href

    if not re.match(r"^https?://", href, flags=re.I):
        return None

    return href.split("#")[0].strip()


def is_track_url(url: str) -> bool:
    if not url or not url.startswith("https://soundcloud.com/"):
        return False

    path = urlparse(url).path.strip("/")
    parts = [p for p in path.split("/") if p]

    if len(parts) != 2:
        return False

    banned_first = {
        "you", "discover", "stream", "upload", "search",
        "charts", "stations", "genres", "tags",
        "pages", "company", "jobs", "imprint", "terms-of-use",
        "about", "blog", "mobile", "download", "developers",
        "support", "legal", "copyright", "press", "newsroom",
        "advertising", "pro", "for-artists", "community",
        "go", "accounts", "settings"
    }

    banned_second = {
        "likes", "followers", "following", "library", "tracks",
        "albums", "reposts", "spotlight", "sets",
        "privacy", "cookies", "newsroom", "getheard",
        "terms", "about", "jobs", "developers", "blog",
        "mobile", "download", "legal", "copyright"
    }

    first = parts[0].lower()
    second = parts[1].lower()

    if first in banned_first or second in banned_second:
        return False

    return True


def build_likes_url(profile_or_likes_url: str) -> str:
    likes_url = profile_or_likes_url.rstrip("/")
    if not likes_url.endswith("/likes"):
        likes_url += "/likes"
    return likes_url


def get_profile_slug_from_url(profile_or_likes_url: str) -> str:
    normalized = profile_or_likes_url.rstrip("/")
    if normalized.endswith("/likes"):
        normalized = normalized[:-6]

    parsed = urlparse(normalized)
    path = parsed.path.strip("/")
    parts = [p for p in path.split("/") if p]

    if not parts:
        return "soundcloud_profile"

    return slugify_profile_name(parts[0])


def build_output_paths(profile_or_likes_url: str):
    profile_slug = get_profile_slug_from_url(profile_or_likes_url)
    csv_file = f"{profile_slug}_likes.csv"
    state_file = f"{profile_slug}_state.json"
    return profile_slug, csv_file, state_file


def classify_link(url: str):
    if not url:
        return None

    try:
        parsed = urlparse(url)
        host = (parsed.netloc or "").lower()
        path = (parsed.path or "").lower()
    except Exception:
        return "other_external"

    if "soundcloud.com" in host:
        if "/download" in path:
            return "soundcloud_download"
        return "soundcloud"

    if "beatport.com" in host:
        return "beatport"

    if "bandcamp.com" in host:
        return "bandcamp"

    if "hypeddit.com" in host:
        return "hypeddit_gate"

    if "toneden.io" in host or "toneden.com" in host:
        return "toneden_gate"

    if re.search(r"\.(mp3|wav|flac|aiff|zip)(\?|$)", url, flags=re.I):
        return "direct_file"

    return "other_external"


def ensure_csv_exists(csv_path: str):
    if os.path.exists(csv_path):
        return

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        writer.writeheader()


def load_rows(csv_path: str):
    rows = []
    if not os.path.exists(csv_path):
        return rows

    with open(csv_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(dict(row))
    return rows


def write_rows(csv_path: str, rows):
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def append_rows(csv_path: str, rows):
    if not rows:
        return
    with open(csv_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        writer.writerows(rows)


def load_existing_track_urls(csv_path: str):
    seen = set()
    if not os.path.exists(csv_path):
        return seen

    try:
        with open(csv_path, "r", newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                url = normalize_url(row.get("track_url", ""))
                if url:
                    seen.add(url)
    except Exception as e:
        print(f"Warning: could not fully read existing CSV '{csv_path}': {e}")

    return seen


def load_state(state_path: str):
    if not os.path.exists(state_path):
        return {}
    try:
        with open(state_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            if isinstance(data, dict):
                return data
    except Exception as e:
        print(f"Warning: could not read state file '{state_path}': {e}")
    return {}


def save_state(state_path: str, state: dict):
    tmp_path = state_path + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(state, f, indent=2)
    os.replace(tmp_path, state_path)


def prompt_for_max_new_urls():
    while True:
        raw = input("Enter max number of NEW tracks to collect this run (leave blank for no limit): ").strip()
        if raw == "":
            return None
        try:
            value = int(raw)
            if value <= 0:
                print("Please enter a positive whole number, or leave blank for no limit.")
                continue
            return value
        except ValueError:
            print("Invalid input. Enter a positive whole number, or leave blank for no limit.")


def prompt_for_gate_limit(default=DEFAULT_GATE_LIMIT):
    raw = input(f"Enter max number of gated links to process this run (blank for {default}): ").strip()
    if not raw:
        return default
    try:
        value = int(raw)
        return max(1, value)
    except ValueError:
        print(f"Invalid input. Using default {default}.")
        return default


# -----------------------------
# Scraper logic
# -----------------------------
def extract_visible_track_data(page):
    js = r"""
    () => {
        const out = [];
        const globalSeen = new Set();

        const bannedFirst = new Set([
            'you', 'discover', 'stream', 'upload', 'search',
            'charts', 'stations', 'genres', 'tags',
            'pages', 'company', 'jobs', 'imprint', 'terms-of-use',
            'about', 'blog', 'mobile', 'download', 'developers',
            'support', 'legal', 'copyright', 'press', 'newsroom',
            'advertising', 'pro', 'for-artists', 'community',
            'go', 'accounts', 'settings'
        ]);

        const bannedSecond = new Set([
            'likes', 'followers', 'following', 'library', 'tracks',
            'albums', 'reposts', 'spotlight', 'sets',
            'privacy', 'cookies', 'newsroom', 'getheard',
            'terms', 'about', 'jobs', 'developers', 'blog',
            'mobile', 'download', 'legal', 'copyright'
        ]);

        function normalizeScUrl(href) {
            if (!href) return null;
            href = href.trim();
            if (!href) return null;
            if (href.startsWith('/')) href = 'https://soundcloud.com' + href;
            if (!href.startsWith('https://soundcloud.com/')) return null;
            return href.split('?')[0].split('#')[0].replace(/\/$/, '');
        }

        function normalizeAnyUrl(href) {
            if (!href) return null;
            href = href.trim();
            if (!href) return null;
            if (href.startsWith('//')) href = 'https:' + href;
            else if (href.startsWith('/')) href = 'https://soundcloud.com' + href;
            if (!/^https?:\/\//i.test(href)) return null;
            return href.split('#')[0];
        }

        function isTrackUrl(url) {
            if (!url || !url.startsWith('https://soundcloud.com/')) return false;
            const path = url.replace('https://soundcloud.com/', '').replace(/^\/+|\/+$/g, '');
            const parts = path.split('/').filter(Boolean);
            if (parts.length !== 2) return false;
            const first = parts[0].toLowerCase();
            const second = parts[1].toLowerCase();
            return !bannedFirst.has(first) && !bannedSecond.has(second);
        }

        function text(el) {
            if (!el) return null;
            const t = (el.innerText || el.textContent || '').replace(/\s+/g, ' ').trim();
            return t || null;
        }

        function attr(el, name) {
            if (!el) return null;
            const v = el.getAttribute(name);
            return v ? v.trim() : null;
        }

        function pickFirstText(root, selectors) {
            for (const sel of selectors) {
                const el = root.querySelector(sel);
                const t = text(el);
                if (t) return t;
            }
            return null;
        }

        function looksLikeGenre(value) {
            if (!value) return false;
            const v = value.trim().toLowerCase();
            if (!v) return false;
            if (v.includes('like') || v.includes('repost') || v.includes('comment') || v.includes('play')) return false;
            if (v.length > 40) return false;
            return true;
        }

        function findGenre(row) {
            const selectors = [
                'a[href*="/tags/"]',
                '.sc-tag',
                '.soundContext__tags a',
                '.soundTitle__tag',
                '.sound__tag',
                '[class*="tag"] a',
                '[class*="genre"]',
                'a[title^="#"]'
            ];

            for (const sel of selectors) {
                for (const el of row.querySelectorAll(sel)) {
                    let t = text(el) || attr(el, 'title') || attr(el, 'aria-label');
                    if (!t) continue;
                    t = t.replace(/^#/, '').trim();
                    if (looksLikeGenre(t)) return t;
                }
            }
            return null;
        }

        function findTrackUrl(row) {
            const anchors = row.querySelectorAll('a[href]');
            for (const a of anchors) {
                const href = normalizeScUrl(a.getAttribute('href') || '');
                if (href && isTrackUrl(href)) return href;
            }
            return null;
        }

        function findTitle(row, trackUrl) {
            const direct = pickFirstText(row, [
                '.soundTitle__title',
                '.soundTitle__titleLink',
                'a.soundTitle__title',
                'a[href*="/"][itemprop="url"]',
            ]);
            if (direct) return direct;

            if (trackUrl) {
                const anchors = row.querySelectorAll('a[href]');
                for (const a of anchors) {
                    const href = normalizeScUrl(a.getAttribute('href') || '');
                    if (href === trackUrl) {
                        const t = text(a) || attr(a, 'title') || attr(a, 'aria-label');
                        if (t) return t;
                    }
                }
            }
            return null;
        }

        function findArtist(row, trackUrl) {
            const direct = pickFirstText(row, [
                '.soundTitle__username',
                '.soundTitle__usernameHero',
                '.soundTitle__username a',
                'a[href^="/"][class*="user"]',
            ]);
            if (direct) return direct;

            const anchors = row.querySelectorAll('a[href]');
            for (const a of anchors) {
                const raw = a.getAttribute('href') || '';
                if (!raw) continue;
                if (raw.startsWith('/')) {
                    const full = 'https://soundcloud.com' + raw;
                    const path = full.replace('https://soundcloud.com/', '').replace(/^\/+|\/+$/g, '');
                    const parts = path.split('/').filter(Boolean);
                    if (parts.length === 1) {
                        const t = text(a) || attr(a, 'title') || attr(a, 'aria-label');
                        if (t) return t;
                    }
                }
            }

            if (trackUrl) {
                const path = trackUrl.replace('https://soundcloud.com/', '');
                const artistSlug = path.split('/')[0];
                if (artistSlug) return artistSlug;
            }
            return null;
        }

        function detectActionLinks(row) {
            const links = row.querySelectorAll('a[href], button, [role="button"]');
            let downloadAvailable = false;
            let downloadTargetUrl = null;
            let buyUrl = null;

            for (const el of links) {
                const label = [
                    text(el),
                    attr(el, 'title'),
                    attr(el, 'aria-label'),
                    attr(el, 'data-testid')
                ].filter(Boolean).join(' ').toLowerCase();

                let href = null;
                if (el.tagName.toLowerCase() === 'a') {
                    href = normalizeAnyUrl(el.getAttribute('href') || '');
                }

                const looksDownload = /free download|download/i.test(label) || (!!href && /download/i.test(href));
                const looksBuy = /\bbuy\b|\bbeatport\b/i.test(label) || (!!href && /beatport\.com/i.test(href));

                if (looksDownload) {
                    downloadAvailable = true;
                    if (!downloadTargetUrl && href) downloadTargetUrl = href;
                }
                if (looksBuy) {
                    if (!buyUrl && href) buyUrl = href;
                }
            }

            return {
                download_available: downloadAvailable,
                download_target_url: downloadTargetUrl,
                buy_url: buyUrl
            };
        }

        const rowSelectors = ['main li', 'main article', '.soundList > li', '.lazyLoadingList__list > li'];
        for (const selector of rowSelectors) {
            const rows = document.querySelectorAll(selector);
            if (!rows.length) continue;

            for (const row of rows) {
                const trackUrl = findTrackUrl(row);
                if (!trackUrl || globalSeen.has(trackUrl)) continue;
                globalSeen.add(trackUrl);

                const title = findTitle(row, trackUrl);
                const artist = findArtist(row, trackUrl);
                const genre = findGenre(row);
                const actions = detectActionLinks(row);

                out.push({
                    track_url: trackUrl,
                    title,
                    artist,
                    genre,
                    download_available: !!actions.download_available,
                    download_target_url: actions.download_target_url || null,
                    buy_url: actions.buy_url || null
                });
            }
            if (out.length > 0) break;
        }
        return out;
    }
    """
    try:
        results = page.evaluate(js)
        return results if isinstance(results, list) else []
    except Exception as e:
        print(f"Warning: extract_visible_track_data failed: {e}")
        return []


def build_csv_row(track: dict, likes_url: str):
    track_url = normalize_url(track.get("track_url"))
    if not track_url or not is_track_url(track_url):
        return None

    download_target_url = normalize_any_url(track.get("download_target_url"))
    buy_url = normalize_any_url(track.get("buy_url"))

    return {
        "track_url": track_url,
        "title": clean_text(track.get("title")),
        "artist": clean_text(track.get("artist")),
        "genre": clean_text(track.get("genre")),
        "download_available": bool(track.get("download_available")),
        "download_target_url": download_target_url,
        "download_target_type": classify_link(download_target_url),
        "buy_url": buy_url,
        "buy_target_type": classify_link(buy_url),
        "discovered_at": utc_now_iso(),
        "source_profile": likes_url,
        "gate_attempted_at": "",
        "gate_status": "",
        "gate_notes": "",
        "final_download_url": "",
    }


def make_state_dict(profile_slug, likes_url, csv_file, seen_urls_count, total_new_this_run, max_new_urls, complete):
    return {
        "profile_slug": profile_slug,
        "profile_url": likes_url,
        "csv_file": csv_file,
        "saved_url_count": seen_urls_count,
        "new_urls_added_this_run": total_new_this_run,
        "max_new_urls_this_run": max_new_urls,
        "last_run_utc": utc_now_iso(),
        "complete": complete,
    }


def scrape_all_likes(profile_or_likes_url: str):
    likes_url = build_likes_url(profile_or_likes_url)
    profile_slug, csv_file, state_file = build_output_paths(profile_or_likes_url)

    if not os.path.exists(BRAVE_PATH):
        print(f"Brave not found at: {BRAVE_PATH}")
        sys.exit(1)

    ensure_csv_exists(csv_file)
    seen_urls = load_existing_track_urls(csv_file)
    previous_state = load_state(state_file)
    max_new_urls = prompt_for_max_new_urls()

    print("\n========== STARTUP SUMMARY ==========")
    print(f"Profile slug: {profile_slug}")
    print(f"Likes URL: {likes_url}")
    print(f"CSV file: {csv_file}")
    print(f"State file: {state_file}")
    print(f"Already saved in CSV: {len(seen_urls)}")
    print(f"Max NEW tracks this run: {'No limit' if max_new_urls is None else max_new_urls}")
    if previous_state:
        print(f"Previous recorded saved_url_count: {previous_state.get('saved_url_count', 'N/A')}")
        print(f"Previous run timestamp: {previous_state.get('last_run_utc', 'N/A')}")
        print(f"Previous completion flag: {previous_state.get('complete', False)}")
    else:
        print("Previous state: none")
    print("====================================\n")

    pending_rows = []
    total_new_this_run = 0
    idle_rounds = 0
    cycle = 0
    scrape_complete = False
    previous_candidate_count = 0
    duplicate_only_rounds = 0

    with sync_playwright() as p:
        browser = p.chromium.launch(
            executable_path=BRAVE_PATH,
            headless=False,
            args=["--incognito"]
        )

        context = browser.new_context()
        page = context.new_page()

        print(f"Opening Brave incognito at: {likes_url}")
        page.goto(likes_url, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
        page.wait_for_timeout(5000)

        while True:
            cycle += 1
            candidate_tracks = extract_visible_track_data(page)
            current_candidate_count = len(candidate_tracks)
            cycle_new = 0

            for track in candidate_tracks:
                track_url = normalize_url(track.get("track_url"))
                if not track_url or not is_track_url(track_url):
                    continue
                if track_url in seen_urls:
                    continue

                row = build_csv_row(track, likes_url)
                if not row:
                    continue

                seen_urls.add(track_url)
                cycle_new += 1
                total_new_this_run += 1
                pending_rows.append(row)

                if max_new_urls is not None and total_new_this_run >= max_new_urls:
                    print(f"Reached requested limit of {max_new_urls} new tracks for this run.")
                    break

            print(
                f"Cycle {cycle}: found {current_candidate_count} candidates, {cycle_new} new, "
                f"pending {len(pending_rows)}, total new this run {total_new_this_run}, grand total known {len(seen_urls)}"
            )

            if len(pending_rows) >= BATCH_SIZE or (max_new_urls is not None and total_new_this_run >= max_new_urls):
                append_rows(csv_file, pending_rows)
                print(f"Saved batch of {len(pending_rows)} rows to {csv_file}")
                pending_rows.clear()
                save_state(
                    state_file,
                    make_state_dict(profile_slug, likes_url, csv_file, len(seen_urls), total_new_this_run, max_new_urls, False)
                )

            if max_new_urls is not None and total_new_this_run >= max_new_urls:
                break

            if current_candidate_count == previous_candidate_count:
                idle_rounds += 1
            else:
                idle_rounds = 0
                previous_candidate_count = current_candidate_count

            if cycle_new == 0 and current_candidate_count > 0:
                duplicate_only_rounds += 1
            else:
                duplicate_only_rounds = 0

            if idle_rounds >= MAX_IDLE_ROUNDS:
                print("Reached idle limit. Page is no longer loading new tracks.")
                scrape_complete = True
                break

            if duplicate_only_rounds >= DUPLICATE_ONLY_THRESHOLD:
                scroll_pixels = FAST_SCROLL_PIXELS
                wait_ms = FAST_SCROLL_WAIT_MS
                mode = "FAST"
            else:
                scroll_pixels = NORMAL_SCROLL_PIXELS
                wait_ms = WAIT_MS
                mode = "NORMAL"

            print(f"Scroll mode: {mode}, duplicate_only_rounds={duplicate_only_rounds}, scroll_pixels={scroll_pixels}, wait_ms={wait_ms}")
            page.mouse.wheel(0, scroll_pixels)
            page.wait_for_timeout(wait_ms)

        if pending_rows:
            append_rows(csv_file, pending_rows)
            print(f"Saved final batch of {len(pending_rows)} rows to {csv_file}")

        save_state(
            state_file,
            make_state_dict(profile_slug, likes_url, csv_file, len(seen_urls), total_new_this_run, max_new_urls, scrape_complete)
        )

        print("\n============= SCRAPE COMPLETE =============")
        print(f"CSV file: {csv_file}")
        print(f"State file: {state_file}")
        print(f"New tracks added this run: {total_new_this_run}")
        print(f"Total tracks now stored for this profile: {len(seen_urls)}")
        print(f"Marked complete: {scrape_complete}")
        print("===========================================\n")

        export_organized_outputs(csv_file)

        browser.close()


# -----------------------------
# Gate automation
# -----------------------------
def wait_for_any_selector(page, selectors, timeout=10000):
    deadline = time.time() + (timeout / 1000)
    while time.time() < deadline:
        for selector in selectors:
            try:
                loc = page.locator(selector)
                if loc.count() > 0 and loc.first.is_visible():
                    return selector
            except Exception:
                pass
        page.wait_for_timeout(250)
    return None


def click_first_matching(page, selectors, timeout=7000):
    matched = wait_for_any_selector(page, selectors, timeout=timeout)
    if not matched:
        return False, None

    try:
        locator = page.locator(matched).first
        locator.scroll_into_view_if_needed()
        locator.click(timeout=3000)
        return True, matched
    except Exception:
        return False, matched


def maybe_accept_cookies(page):
    selectors = [
        'button:has-text("Accept")',
        'button:has-text("I agree")',
        'button:has-text("Allow all")',
        'button:has-text("Got it")',
        '[id*="cookie"] button',
        '[class*="cookie"] button',
    ]
    click_first_matching(page, selectors, timeout=2500)


def capture_final_download_url(page):
    selectors = [
        'a[download]',
        'a[href*="download"]',
        'a[href$=".zip"]',
        'a[href$=".mp3"]',
        'a[href$=".wav"]',
        'a[href$=".flac"]',
        'a[href$=".aiff"]',
    ]
    for selector in selectors:
        try:
            loc = page.locator(selector)
            if loc.count() > 0:
                href = loc.first.get_attribute('href')
                href = normalize_any_url(href)
                if href:
                    return href
        except Exception:
            pass
    return None


def handle_soundcloud_oauth_popup(context, parent_page, notes):
    popup = None
    try:
        popup = context.wait_for_event("page", timeout=10000)
    except PlaywrightTimeoutError:
        return False, notes + ["No popup appeared for SoundCloud connect."]

    try:
        popup.wait_for_load_state("domcontentloaded", timeout=15000)
        maybe_accept_cookies(popup)

        # Best case: session already logged in and approval page is shown.
        approve_selectors = [
            'button:has-text("Connect")',
            'button:has-text("Authorize")',
            'button:has-text("Allow")',
            'input[type="submit"][value*="Connect"]',
            'input[type="submit"][value*="Authorize"]',
        ]
        clicked, which = click_first_matching(popup, approve_selectors, timeout=12000)
        if clicked:
            notes.append(f"Clicked OAuth approval control: {which}")
        else:
            notes.append("OAuth popup opened. No approval button was auto-clicked; session may need manual login/approval.")

        try:
            popup.wait_for_close(timeout=20000)
        except PlaywrightTimeoutError:
            notes.append("OAuth popup remained open after attempt.")

        parent_page.wait_for_timeout(4000)
        return True, notes
    except Exception as e:
        notes.append(f"OAuth popup handling error: {e}")
        return False, notes


def process_hypeddit_gate(context, gate_page):
    notes = []
    gate_page.wait_for_load_state("domcontentloaded", timeout=PAGE_TIMEOUT_MS)
    maybe_accept_cookies(gate_page)
    gate_page.wait_for_timeout(2500)

    connect_selectors = [
        'button:has-text("Connect SoundCloud")',
        'a:has-text("Connect SoundCloud")',
        'button:has-text("Connect")',
        'a:has-text("Connect")',
        '[href*="soundcloud.com/connect"]',
        '[href*="soundcloud.com/oauth"]',
        '[data-testid*="soundcloud"]',
    ]
    clicked, which = click_first_matching(gate_page, connect_selectors, timeout=12000)
    if clicked:
        notes.append(f"Clicked Hypeedit connect control: {which}")
        _, notes = handle_soundcloud_oauth_popup(context, gate_page, notes)
    else:
        notes.append("No obvious Hypeedit connect control found.")

    unlock_selectors = [
        'button:has-text("Unlock")',
        'a:has-text("Unlock")',
        'button:has-text("Download")',
        'a:has-text("Download")',
        'button:has-text("Get Download")',
        'a:has-text("Get Download")',
    ]
    clicked, which = click_first_matching(gate_page, unlock_selectors, timeout=10000)
    if clicked:
        notes.append(f"Clicked Hypeedit unlock/download control: {which}")
    else:
        notes.append("No Hypeedit unlock/download button was auto-clicked.")

    gate_page.wait_for_timeout(3000)
    final_url = capture_final_download_url(gate_page)
    return {
        "status": "attempted",
        "notes": " | ".join(notes),
        "final_download_url": final_url or "",
    }


def process_toneden_gate(context, gate_page):
    notes = []
    gate_page.wait_for_load_state("domcontentloaded", timeout=PAGE_TIMEOUT_MS)
    maybe_accept_cookies(gate_page)
    gate_page.wait_for_timeout(2500)

    connect_selectors = [
        'button:has-text("Connect SoundCloud")',
        'a:has-text("Connect SoundCloud")',
        'button:has-text("Continue with SoundCloud")',
        'a:has-text("Continue with SoundCloud")',
        'button:has-text("Connect")',
        '[href*="soundcloud.com/connect"]',
        '[href*="soundcloud.com/oauth"]',
    ]
    clicked, which = click_first_matching(gate_page, connect_selectors, timeout=12000)
    if clicked:
        notes.append(f"Clicked ToneDen connect control: {which}")
        _, notes = handle_soundcloud_oauth_popup(context, gate_page, notes)
    else:
        notes.append("No obvious ToneDen connect control found.")

    unlock_selectors = [
        'button:has-text("Unlock")',
        'a:has-text("Unlock")',
        'button:has-text("Download")',
        'a:has-text("Download")',
        'button:has-text("Get it")',
        'a:has-text("Get it")',
    ]
    clicked, which = click_first_matching(gate_page, unlock_selectors, timeout=10000)
    if clicked:
        notes.append(f"Clicked ToneDen unlock/download control: {which}")
    else:
        notes.append("No ToneDen unlock/download button was auto-clicked.")

    gate_page.wait_for_timeout(3000)
    final_url = capture_final_download_url(gate_page)
    return {
        "status": "attempted",
        "notes": " | ".join(notes),
        "final_download_url": final_url or "",
    }


def process_generic_gate(context, gate_page):
    notes = []
    gate_page.wait_for_load_state("domcontentloaded", timeout=PAGE_TIMEOUT_MS)
    maybe_accept_cookies(gate_page)
    gate_page.wait_for_timeout(2500)

    form_fill_rules = [
        ('input[type="email"]', os.getenv('GATE_EMAIL', '')),
        ('input[name*="email" i]', os.getenv('GATE_EMAIL', '')),
        ('input[placeholder*="email" i]', os.getenv('GATE_EMAIL', '')),
        ('input[name*="first" i]', os.getenv('GATE_FIRST_NAME', '')),
        ('input[name*="last" i]', os.getenv('GATE_LAST_NAME', '')),
        ('input[name*="name" i]:not([type="hidden"])', os.getenv('GATE_NAME', '')),
    ]

    for selector, value in form_fill_rules:
        if not value:
            continue
        try:
            loc = gate_page.locator(selector)
            if loc.count() > 0 and loc.first.is_visible():
                loc.first.fill(value, timeout=2000)
                notes.append(f"Filled field: {selector}")
        except Exception:
            pass

    connect_selectors = [
        'button:has-text("Connect SoundCloud")',
        'a:has-text("Connect SoundCloud")',
        'button:has-text("Continue with SoundCloud")',
        'a:has-text("Continue with SoundCloud")',
        '[href*="soundcloud.com/connect"]',
        '[href*="soundcloud.com/oauth"]',
    ]
    clicked, which = click_first_matching(gate_page, connect_selectors, timeout=8000)
    if clicked:
        notes.append(f"Clicked generic SoundCloud connect control: {which}")
        _, notes = handle_soundcloud_oauth_popup(context, gate_page, notes)

    submit_selectors = [
        'button:has-text("Submit")',
        'button:has-text("Continue")',
        'button:has-text("Unlock")',
        'button:has-text("Download")',
        'a:has-text("Download")',
        'input[type="submit"]',
    ]
    clicked, which = click_first_matching(gate_page, submit_selectors, timeout=10000)
    if clicked:
        notes.append(f"Clicked generic submit/download control: {which}")
    else:
        notes.append("No generic form submit or download control was auto-clicked.")

    gate_page.wait_for_timeout(3000)
    final_url = capture_final_download_url(gate_page)
    return {
        "status": "attempted",
        "notes": " | ".join(notes),
        "final_download_url": final_url or "",
    }


def process_gate_url(context, row):
    target_url = normalize_any_url(row.get("download_target_url"))
    if not target_url:
        return {
            "status": "skipped",
            "notes": "No download target URL present.",
            "final_download_url": "",
        }

    target_type = classify_link(target_url)
    if target_type in {"beatport", "bandcamp", "soundcloud", "other_external"}:
        return {
            "status": "skipped",
            "notes": f"Unsupported gate type for automation: {target_type}",
            "final_download_url": "",
        }

    page = context.new_page()
    try:
        print(f"Opening gate URL: {target_url}")
        page.goto(target_url, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
        page.wait_for_timeout(2500)

        if target_type == "hypeddit_gate":
            result = process_hypeddit_gate(context, page)
        elif target_type == "toneden_gate":
            result = process_toneden_gate(context, page)
        elif target_type == "direct_file":
            result = {
                "status": "ready",
                "notes": "Direct file link already present.",
                "final_download_url": target_url,
            }
        elif target_type == "soundcloud_download":
            result = {
                "status": "ready",
                "notes": "SoundCloud download link already present.",
                "final_download_url": target_url,
            }
        else:
            result = process_generic_gate(context, page)

        return result
    except Exception as e:
        return {
            "status": "error",
            "notes": f"Gate processing failed: {e}",
            "final_download_url": "",
        }
    finally:
        try:
            page.close()
        except Exception:
            pass


def process_csv_gates(csv_path: str):
    rows = load_rows(csv_path)
    if not rows:
        print(f"No rows found in {csv_path}")
        return

    gate_candidates = []
    for i, row in enumerate(rows):
        target_url = normalize_any_url(row.get("download_target_url"))
        if not target_url:
            continue
        if str(row.get("gate_status", "")).strip().lower() in {"ready", "attempted", "success"}:
            continue
        gate_candidates.append((i, row))

    if not gate_candidates:
        print("No unprocessed gated links found in CSV.")
        return

    limit = prompt_for_gate_limit()
    print(f"Found {len(gate_candidates)} unprocessed gate candidates. Processing up to {limit}.")

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            executable_path=BRAVE_PATH,
            headless=False,
            args=["--start-maximized"],
            accept_downloads=True,
        )

        print("\nUsing a persistent browser profile for gate automation.")
        print(f"Profile dir: {USER_DATA_DIR}")
        print("Sign into SoundCloud manually once in this browser profile if OAuth approval is not automatic.\n")

        processed = 0
        for index, row in gate_candidates:
            if processed >= limit:
                break

            result = process_gate_url(context, row)
            rows[index]["gate_attempted_at"] = utc_now_iso()
            rows[index]["gate_status"] = result.get("status", "attempted")
            rows[index]["gate_notes"] = result.get("notes", "")
            rows[index]["final_download_url"] = result.get("final_download_url", "")

            write_rows(csv_path, rows)
            processed += 1
            print(
                f"[{processed}/{limit}] {row.get('track_url')} -> {rows[index]['gate_status']} | {rows[index]['gate_notes']}"
            )

        context.close()

    print("\n============= GATE RUN COMPLETE =============")
    print(f"CSV file updated: {csv_path}")
    print(f"Rows processed this run: {processed}")
    print("============================================\n")

    export_organized_outputs(csv_path)



# -----------------------------
# Organized CSV / Excel outputs
# -----------------------------
ORGANIZED_EXTRA_FIELDS = [
    "original_genre",
    "genre_rule_applied",
    "genre_sheet",
    "download_website",
    "buy_website",
]

UNCLEAR_GENRE = "Unclear / Needs Review"

# Add or edit these rules whenever you want a different genre layout.
# Matching is intentionally forgiving about spaces, hyphens, slashes, and capitalization.
GENRE_RULES = [
    ("Hip-Hop / Rap", [
        "hip hop", "hip-hop", "hiphop", "rap", "hip hop rap", "hiphop rap",
        "rap hip hop", "rap/hip hop", "hiphop/rap", "hip-hop & rap", "hip hop-rap",
    ]),
    ("R&B / Soul", [
        "r&b", "r & b", "rnb", "rhythm and blues", "soul", "r&b soul",
        "r&b & soul", "r&b-soul", "neo soul", "r&b soul neo soul",
    ]),
    ("Afro House", ["afro house", "afrohouse"]),
    ("Minimal / Deep Tech", [
        "minimal deep tech", "minimal-deep-tech", "minimal - deep tech",
        "minimal/deep tech", "minimal deepteech", "minimal - deepteech",
        "minimal deep tech 1",
    ]),
    ("Deep Tech / Tech House", [
        "deep tech", "deep-tech", "tech house deep", "tech house - deep",
        "deep tech house", "deep house tech", "tech-house-deep",
    ]),
    ("Tech House", ["tech house", "tech-house", "techhouse"]),
    ("Deep House", ["deep house", "deephouse", "deep-house"]),
    ("House", ["house", "housedance", "club house"]),
    ("Dance / EDM", ["dance", "edm", "dance edm", "dance & edm", "electronic dance", "electronic-dance"]),
    ("Electronic", ["electronic", "electronica", "electronica/dance", "electro"]),
    ("Reggae", ["reggae", "reggae 1"]),
    ("Dancehall", ["dancehall", "dance hall"]),
    ("Latin Urban", ["latin urban", "latin - urban", "latin-urban"]),
    ("Christian / Gospel", ["gospel", "christian", "gospel christian", "gospel-christian", "christian gospel", "christian & gospel"]),
    ("Pop", ["pop", "popular"]),
    ("Trap", ["trap"]),
    ("Drum & Bass", ["drum and bass", "drum & bass", "dnb", "d&b"]),
    ("Dubstep", ["dubstep", "dub step"]),
    ("Amapiano", ["amapiano"]),
    ("Afrobeats", ["afrobeats", "afro beats", "afrobeat"]),
    ("Jersey Club", ["jersey club", "jerseyclub"]),
    ("Phonk", ["phonk"]),
]

HOUSE_HINTS = [
    "house", "hypeddit", "hypeedit", "tech house", "deep tech", "deep house",
    "afro house", "minimal", "jackin", "garage", "ukg"
]

NON_GENRE_HINTS = [
    "records", "recordings", "music group", "label", "collective", "official",
    "playlist", "premiere", "download", "free download", "soundcloud", "unknown",
]


def compact_genre_key(value):
    value = clean_text(value) or ""
    value = value.lower().replace("&", " and ")
    value = re.sub(r"[#_]+", " ", value)
    value = re.sub(r"[\\/|,+:;()\[\]{}]+", " ", value)
    value = re.sub(r"[-]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def proper_case_genre(value):
    value = compact_genre_key(value)
    if not value:
        return UNCLEAR_GENRE
    words = []
    for word in value.split():
        if word in {"rnb", "edm", "dnb"}:
            words.append(word.upper())
        elif word == "and":
            words.append("&")
        else:
            words.append(word.capitalize())
    return " ".join(words)


def looks_unclear_genre(value):
    raw = clean_text(value)
    if not raw:
        return True
    key = compact_genre_key(raw)
    if not key:
        return True
    if len(key) > 45:
        return True
    if any(hint in key for hint in NON_GENRE_HINTS):
        return True
    # Many artist/label names have several words and no recognizable genre hints.
    if len(key.split()) >= 5 and not any(hint in key for hint in HOUSE_HINTS):
        return True
    return False


def website_from_url(url):
    url = normalize_any_url(url)
    if not url:
        return ""
    try:
        host = urlparse(url).netloc.lower().replace("www.", "")
        return host
    except Exception:
        return ""


def standardize_genre(row):
    original = clean_text(row.get("genre")) or clean_text(row.get("original_genre")) or ""
    download_url = " ".join([
        str(row.get("download_target_url", "")),
        str(row.get("download_target_type", "")),
        str(row.get("buy_url", "")),
        str(row.get("buy_target_type", "")),
    ]).lower()

    # Your current organizing rule: Hypeddit/Hypeedit tracks go under House.
    if "hypeddit" in download_url or "hypeedit" in download_url:
        return "House", "Hypeddit/Hypeedit link rule"

    key = compact_genre_key(original)
    if not key:
        return UNCLEAR_GENRE, "Missing genre"

    for standard, variants in GENRE_RULES:
        variant_keys = {compact_genre_key(v) for v in variants}
        if key in variant_keys:
            return standard, "Exact normalized genre match"
        # Match mixed labels like hiphop/rap or dance-edm after compaction.
        if any(vk and vk in key for vk in variant_keys if len(vk) >= 5):
            return standard, "Partial normalized genre match"

    if any(hint in key for hint in HOUSE_HINTS):
        return "House", "House-related keyword"

    if looks_unclear_genre(original):
        return UNCLEAR_GENRE, "Looks like non-genre, typo, blank, artist, label, or playlist"

    return proper_case_genre(original), "Cleaned capitalization/spacing"


def safe_sheet_name(name, used=None):
    used = used if used is not None else set()
    clean = re.sub(r"[\\/*?:\[\]]", "-", str(name)).strip()
    clean = clean[:31] or "Sheet"
    base = clean
    n = 2
    while clean in used:
        suffix = f" {n}"
        clean = (base[:31 - len(suffix)] + suffix).strip()
        n += 1
    used.add(clean)
    return clean


def enrich_rows_for_organized_output(rows):
    out = []
    for row in rows:
        new_row = dict(row)
        original_genre = clean_text(row.get("genre")) or ""
        standard, reason = standardize_genre(row)
        new_row["original_genre"] = original_genre
        new_row["genre"] = standard
        new_row["genre_rule_applied"] = reason
        new_row["genre_sheet"] = standard
        new_row["download_website"] = website_from_url(row.get("download_target_url") or row.get("final_download_url"))
        new_row["buy_website"] = website_from_url(row.get("buy_url"))
        out.append(new_row)
    return out


def dedupe_exact_rows(rows):
    seen = set()
    result = []
    for row in rows:
        key = tuple((str(k), str(row.get(k, ""))) for k in sorted(row.keys(), key=lambda x: "" if x is None else str(x)))
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def write_cleaned_csv(source_csv_path):
    rows = load_rows(source_csv_path)
    cleaned_rows = dedupe_exact_rows(enrich_rows_for_organized_output(rows))
    output_path = source_csv_path.replace(".csv", "_cleaned.csv")
    fieldnames = list(CSV_FIELDS)
    for field in ORGANIZED_EXTRA_FIELDS:
        if field not in fieldnames:
            fieldnames.append(field)
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(cleaned_rows)
    return output_path, cleaned_rows, fieldnames


def write_usb_macro_bas(output_bas_path="AddToUSBDownloadList.bas"):
    macro = r'''Attribute VB_Name = "USBDownloadMacros"
Option Explicit

Sub AddToUSBDownloadList()
    Dim wsSource As Worksheet
    Dim wsUSB As Worksheet
    Dim selectedRange As Range
    Dim rowRange As Range
    Dim nextRow As Long
    Dim urlCol As Long
    Dim sourceUrl As String
    Dim found As Range

    Set wsSource = ActiveSheet

    If TypeName(Selection) <> "Range" Then
        MsgBox "Please select the track rows you want to add.", vbExclamation
        Exit Sub
    End If

    Set selectedRange = Selection

    On Error Resume Next
    Set wsUSB = ThisWorkbook.Worksheets("USB Download List")
    On Error GoTo 0

    If wsUSB Is Nothing Then
        Set wsUSB = ThisWorkbook.Worksheets.Add(After:=ThisWorkbook.Sheets(ThisWorkbook.Sheets.Count))
        wsUSB.Name = "USB Download List"
        wsSource.Rows(1).Copy Destination:=wsUSB.Rows(1)
    End If

    urlCol = 1
    On Error Resume Next
    urlCol = Application.Match("track_url", wsSource.Rows(1), 0)
    On Error GoTo 0
    If urlCol = 0 Then urlCol = 1

    For Each rowRange In selectedRange.Rows
        If rowRange.Row <> 1 Then
            sourceUrl = CStr(wsSource.Cells(rowRange.Row, urlCol).Value)

            Set found = Nothing
            If Len(sourceUrl) > 0 Then
                Set found = wsUSB.Columns(urlCol).Find(What:=sourceUrl, LookIn:=xlValues, LookAt:=xlWhole)
            End If

            If found Is Nothing Then
                nextRow = wsUSB.Cells(wsUSB.Rows.Count, "A").End(xlUp).Row + 1
                wsSource.Rows(rowRange.Row).Copy Destination:=wsUSB.Rows(nextRow)
            End If
        End If
    Next rowRange

    MsgBox "Selected tracks added to USB Download List.", vbInformation
End Sub
'''
    with open(output_bas_path, "w", encoding="utf-8") as f:
        f.write(macro)
    return output_bas_path


def write_organized_workbook(source_csv_path, cleaned_rows, fieldnames):
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        print("openpyxl is not installed. Run: python3 -m pip install openpyxl")
        return None

    output_path = source_csv_path.replace(".csv", "_organized.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "All Tracks Cleaned"

    def write_sheet(sheet, rows):
        sheet.append(fieldnames)
        for r in rows:
            sheet.append([r.get(f, "") for f in fieldnames])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col_idx, field in enumerate(fieldnames, start=1):
            max_len = max([len(str(field))] + [len(str(row.get(field, ""))) for row in rows[:200]])
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 45)

    write_sheet(ws, cleaned_rows)

    mapping = {}
    for row in cleaned_rows:
        orig = row.get("original_genre", "") or "(blank)"
        std = row.get("genre_sheet", UNCLEAR_GENRE)
        reason = row.get("genre_rule_applied", "")
        mapping[(orig, std, reason)] = True

    map_ws = wb.create_sheet("Genre Mapping")
    map_ws.append(["Original Genre / Sheet Name", "Standardized Genre", "Reason"])
    for orig, std, reason in sorted(mapping.keys(), key=lambda x: (x[1], x[0])):
        map_ws.append([orig, std, reason])
    map_ws.freeze_panes = "A2"
    map_ws.auto_filter.ref = map_ws.dimensions
    map_ws.column_dimensions["A"].width = 35
    map_ws.column_dimensions["B"].width = 28
    map_ws.column_dimensions["C"].width = 45

    usb_ws = wb.create_sheet("USB Download List")
    usb_ws.append(fieldnames)
    usb_ws.freeze_panes = "A2"
    usb_ws.auto_filter.ref = usb_ws.dimensions

    by_genre = {}
    for row in cleaned_rows:
        by_genre.setdefault(row.get("genre_sheet", UNCLEAR_GENRE), []).append(row)

    used_sheet_names = {ws.title for ws in wb.worksheets}
    for genre, rows in sorted(by_genre.items(), key=lambda item: item[0].lower()):
        sheet_name = safe_sheet_name(genre, used_sheet_names)
        genre_ws = wb.create_sheet(sheet_name)
        write_sheet(genre_ws, rows)

    wb.save(output_path)
    return output_path


def export_organized_outputs(csv_path):
    if not os.path.exists(csv_path):
        print(f"Cannot export organized outputs. Missing CSV: {csv_path}")
        return
    cleaned_csv_path, cleaned_rows, fieldnames = write_cleaned_csv(csv_path)
    workbook_path = write_organized_workbook(csv_path, cleaned_rows, fieldnames)
    macro_path = write_usb_macro_bas("AddToUSBDownloadList.bas")
    print("\n============= ORGANIZED OUTPUTS CREATED =============")
    print(f"Regular cleaned CSV: {cleaned_csv_path}")
    if workbook_path:
        print(f"Excel workbook with sheets: {workbook_path}")
    print(f"Macro module for Mac Excel: {macro_path}")
    print("Note: CSV files cannot contain sheets or macros. Import the .bas file into a macro-enabled .xlsm workbook, or save the workbook as .xlsm after adding the macro in Excel.")
    print("=====================================================\n")

# -----------------------------
# CLI entrypoint
# -----------------------------
def print_usage():
    print("Usage:")
    print("  python3 sc_gate_automation.py scrape https://soundcloud.com/PROFILE")
    print("  python3 sc_gate_automation.py scrape https://soundcloud.com/PROFILE/likes")
    print("  python3 sc_gate_automation.py gates path/to/profile_likes.csv")
    print("")
    print("After scraping or gate processing, the script also creates:")
    print("  *_cleaned.csv, *_organized.xlsx, and AddToUSBDownloadList.bas")
    print("")
    print("Optional env vars for generic forms:")
    print("  GATE_EMAIL=you@example.com")
    print("  GATE_FIRST_NAME=Marcus")
    print("  GATE_LAST_NAME=Martin")
    print("  GATE_NAME=Marcus Martin")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print_usage()
        sys.exit(1)

    command = sys.argv[1].strip().lower()
    target = sys.argv[2].strip()

    if command == "scrape":
        scrape_all_likes(target)
    elif command == "gates":
        process_csv_gates(target)
    else:
        print_usage()
        sys.exit(1)
